"""tools/case_folder.py on small generated documents. No model, platform or network is used.

Word and Excel files are made with python-docx and openpyxl, so those tests skip where the libraries are missing.
"""
import datetime
import importlib.util
import json
import shutil
import struct
import subprocess
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path
from unittest.mock import patch

from tools import case_folder

HAS = {name: importlib.util.find_spec(name) is not None for name in ('docx', 'openpyxl', 'pypdf')}


def png(path, w=4, h=3):
    raw = b''.join(b'\x00' + b'\xff\x00\x00' * w for _ in range(h))
    chunk = lambda t, d: struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    path.write_bytes(b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0))
                     + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def pdf(path, text):
    stream = f'BT /F1 12 Tf 20 100 Td ({text}) Tj ET'
    objs = ['<< /Type /Catalog /Pages 2 0 R >>', '<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
            '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
            f'<< /Length {len(stream)} >>\nstream\n{stream}\nendstream', '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    body, offsets = b'%PDF-1.4\n', []
    for i, o in enumerate(objs, 1):
        offsets.append(len(body))
        body += f'{i} 0 obj\n{o}\nendobj\n'.encode()
    xref = len(body)
    body += f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode() + b''.join(f'{o:010d} 00000 n \n'.encode() for o in offsets)
    path.write_bytes(body + f'trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())


def pdf_reader_available():
    return HAS['pypdf'] or shutil.which('pdftotext') is not None


class CaseFolderTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.case = self.root / 'case'
        self.case.mkdir()
        (self.case / 'BRIEF.md').write_text('# Brief\n\nStart here.\n')
        (self.case / 'README.md').write_text('# Documents\n\n- `flow.png`: the call flow | drawing\n- `BRIEF.md`: the start\n'
                                             '- `customers.csv`: the test customers\n')
        png(self.case / 'flow.png')
        (self.case / 'customers.csv').write_bytes('ad;soyad\nŞule;Işık\n'.encode('cp1254'))
        (self.case / 'examples').mkdir()
        (self.case / 'examples' / 'call-1.md').write_text('Agent: Merhaba.\n')
        (self.case / 'blob.bin').write_bytes(bytes(range(256)))
        (self.case / '.DS_Store').write_bytes(b'\x00')

    def bundle(self, source=None):
        out = self.root / 'bundle'
        return case_folder.convert(source or self.case, out), out

    def test_order_index_and_text_files(self):
        entries, out = self.bundle()
        self.assertEqual([e['source'] for e in entries],
                         ['BRIEF.md', 'README.md', 'flow.png', 'customers.csv', 'blob.bin', 'examples/call-1.md'])
        by = {e['source']: e for e in entries}
        self.assertEqual(by['flow.png']['readme'], '`flow.png`: the call flow | drawing')
        self.assertEqual(by['flow.png']['images'], ['files/flow.png'])
        self.assertEqual((out / 'files/flow.png').read_bytes(), (self.case / 'flow.png').read_bytes())
        self.assertEqual((out / by['customers.csv']['text'][0]).read_text(), 'ad;soyad\nŞule;Işık\n')
        self.assertEqual(by['blob.bin']['kind'], 'binary')
        self.assertFalse((out / 'files/.DS_Store').exists())
        index = (out / 'INDEX.md').read_text()
        self.assertLess(index.index('| 1 | BRIEF.md |'), index.index('| 2 | README.md |'))
        self.assertIn('the call flow \\| drawing', index)
        self.assertIn('- files/flow.png', index)
        self.assertIn('blob.bin: binary file with no converter', index)
        self.assertEqual(json.loads((out / 'manifest.json').read_text()), entries)

    def test_large_text_file_is_marked_in_the_index(self):
        (self.case / 'big.json').write_text(json.dumps([{'n': i, 'pad': 'x' * 100} for i in range(1500)]))
        entries, out = self.bundle()
        size = (out / 'text/big.json').stat().st_size // 1000
        self.assertGreater(size, 100)
        index = (out / 'INDEX.md').read_text()
        self.assertIn(f'| text/big.json (LARGE, {size} KB: query it with python3 -I, json or csv, instead of reading it whole) |', index)
        self.assertEqual(index.count('LARGE'), 1)
        self.assertEqual({e['source']: e['text'] for e in entries}['big.json'], ['text/big.json'])

    def test_zip_with_one_top_folder(self):
        archive = self.root / 'case.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            for p in self.case.rglob('*'):
                z.write(p, 'Final Case/' + p.relative_to(self.case).as_posix())
            z.writestr('__MACOSX/Final Case/._BRIEF.md', b'\x00')
        entries, out = self.bundle(archive)
        self.assertEqual(entries[0]['source'], 'BRIEF.md')
        self.assertNotIn('__MACOSX', ' '.join(e['source'] for e in entries))
        self.assertIn('Converted from `case.zip`', (out / 'INDEX.md').read_text())

    @unittest.skipUnless(HAS['docx'], 'python-docx is not installed')
    def test_word_text_table_and_picture(self):
        import docx
        d = docx.Document()
        d.add_paragraph('POST /verify_identity checks the caller.')
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = 'customer_id', 'string'
        d.add_picture(str(self.case / 'flow.png'))
        d.save(self.case / 'contract.docx')
        entries, out = self.bundle()
        e = next(e for e in entries if e['source'] == 'contract.docx')
        text = (out / e['text'][0]).read_text()
        self.assertIsNone(e['error'])
        for part in ('POST /verify_identity checks the caller.', 'customer_id', 'string'):
            self.assertIn(part, text)
        if shutil.which('pandoc'):
            self.assertEqual(len(e['images']), 1)
            self.assertTrue((out / e['images'][0]).is_file())
        with patch.object(case_folder, 'run', side_effect=FileNotFoundError('no pandoc')):   # the python-docx fallback
            entries, out = self.bundle()
        e = next(e for e in entries if e['source'] == 'contract.docx')
        self.assertIn('customer_id | string', (out / e['text'][0]).read_text())

    @unittest.skipUnless(HAS['openpyxl'], 'openpyxl is not installed')
    def test_excel_sheets_to_markdown_and_csv(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Customers'
        ws.append(['id', 'born', 'phone_last4', 'rate', 'total', 'note', 'amount'])
        ws.append(['C-1', datetime.date(1990, 5, 17), 471, 0.05, '=D2*100', 'a | b', 420.0])
        ws['C2'].number_format, ws['D2'].number_format = '0000', '0%'
        wb.create_sheet('Ürünler (list)').append(['p1', 126.5])
        wb.save(self.case / 'customers.xlsx')
        entries, out = self.bundle()
        e = next(e for e in entries if e['source'] == 'customers.xlsx')
        self.assertEqual(e['text'], ['text/customers.xlsx.md', 'text/customers.xlsx.Customers.csv', 'text/customers.xlsx.Ürünler_list.csv'])
        md = (out / e['text'][0]).read_text()
        self.assertIn('| C-1 | 1990-05-17 | 0471 | 5% | =D2*100 | a \\| b | 420 |', md)
        self.assertIn('## Sheet: Ürünler (list)', md)
        self.assertEqual((out / e['text'][1]).read_text().splitlines()[1], 'C-1,1990-05-17,0471,5%,=D2*100,a | b,420')

    @unittest.skipUnless(pdf_reader_available(), 'neither pdftotext nor pypdf is available')
    def test_pdf_text_with_page_marks(self):
        pdf(self.case / 'qa.pdf', 'Q: Can the agent transfer? A: Yes.')
        entries, out = self.bundle()
        e = next(e for e in entries if e['source'] == 'qa.pdf')
        text = (out / e['text'][0]).read_text()
        self.assertTrue(text.startswith('--- page 1 ---'))
        self.assertIn('Q: Can the agent transfer? A: Yes.', text)

    @unittest.skipUnless(pdf_reader_available(), 'neither pdftotext nor pypdf is available')
    def test_pdf_without_text_points_to_the_original(self):
        pdf(self.case / 'scan.pdf', '')
        entries, out = self.bundle()
        e = next(e for e in entries if e['source'] == 'scan.pdf')
        self.assertIn('almost no text', e['error'])
        self.assertIn('scan.pdf: almost no text', (out / 'INDEX.md').read_text())

    @unittest.skipUnless(HAS['pypdf'], 'pypdf is not installed')
    def test_pdf_falls_back_to_pypdf(self):
        pdf(self.case / 'qa.pdf', 'Fallback text')
        with patch.object(case_folder, 'run', side_effect=subprocess.CalledProcessError(1, 'pdftotext')):
            entries, out = self.bundle()
        e = next(e for e in entries if e['source'] == 'qa.pdf')
        self.assertIn('Fallback text', (out / e['text'][0]).read_text())

    def test_failed_conversion_is_listed_not_fatal(self):
        (self.case / 'broken.docx').write_bytes(b'not a zip')
        with patch.object(case_folder, 'run', side_effect=FileNotFoundError('missing tool')):
            entries, out = self.bundle()
        e = next(e for e in entries if e['source'] == 'broken.docx')
        self.assertTrue(e['error'])
        self.assertTrue((out / 'files/broken.docx').is_file())
        self.assertIn('broken.docx: RuntimeError', (out / 'INDEX.md').read_text())


if __name__ == '__main__':
    unittest.main()
