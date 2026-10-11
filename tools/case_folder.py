#!/usr/bin/env python3
"""Turn the final case folder (or its zip) into a bundle of text the builder can read.

    python3 tools/case_folder.py <case folder or .zip> <bundle dir>

Every original file is copied unchanged to <bundle>/files/. Readable versions go to <bundle>/text/:
Markdown and other text as it is; Word through pandoc (python-docx, then textutil, if that fails);
PDF through pdftotext (pypdf if that fails); Excel through openpyxl as a Markdown table and a CSV per
sheet; legacy .doc/.xls/.ppt through LibreOffice first. Images (the flow drawing, pictures inside Word
files) stay images and are listed so Claude can open them with its Read tool. <bundle>/INDEX.md lists
every document in reading order, BRIEF.md first and README.md second, with the line README.md gives it;
<bundle>/manifest.json holds the same list for programs.
"""
import csv, datetime, io, json, re, shutil, subprocess, sys, tempfile, zipfile
from pathlib import Path

TEXT = {'.md', '.markdown', '.txt', '.csv', '.tsv', '.json', '.yaml', '.yml', '.xml', '.html', '.htm', '.mmd', '.puml'}
WORD = {'.docx', '.odt', '.rtf', '.pptx'}
EXCEL = {'.xlsx', '.xlsm'}
IMAGE = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.svg', '.bmp', '.tif', '.tiff'}
LEGACY = {'.doc': 'docx', '.xls': 'xlsx', '.ppt': 'pptx'}


def run(*cmd, cwd=None):
    return subprocess.run(cmd, capture_output=True, text=True, timeout=180, check=True, cwd=cwd).stdout


def decode(data, any_encoding):
    if data[:2] in (b'\xff\xfe', b'\xfe\xff'):
        return data.decode('utf-16')
    try:
        return data.decode('utf-8-sig')
    except UnicodeDecodeError:
        if not any_encoding:
            raise
        return data.decode('cp1254', errors='replace')   # a CSV saved by Excel in a Windows code page


def word_text(src, media):
    errors = []
    try:   # pandoc keeps tables and headings, and extracts embedded pictures into `media` (linked relative to the text)
        return run('pandoc', str(src), '-t', 'gfm', '--wrap=none', f'--extract-media={media.name}', cwd=media.parent)
    except Exception as e:
        errors.append(f'pandoc: {e}')
    try:
        import docx
        d = docx.Document(str(src))
        rows = [' | '.join(c.text.strip() for c in r.cells) for t in d.tables for r in t.rows]
        return '\n\n'.join([p.text for p in d.paragraphs] + rows)
    except Exception as e:
        errors.append(f'python-docx: {e}')
    try:
        return run('textutil', '-convert', 'txt', '-stdout', str(src))
    except Exception as e:
        errors.append(f'textutil: {e}')
    raise RuntimeError('; '.join(errors))


def pdf_text(src):
    try:
        pages = run('pdftotext', '-layout', str(src), '-').split('\f')
    except Exception as first:
        try:
            import pypdf
            pages = [p.extract_text() or '' for p in pypdf.PdfReader(str(src)).pages]
        except Exception as e:
            raise RuntimeError(f'pdftotext: {first}; pypdf: {e}')
    pages = pages[:-1] if len(pages) > 1 and not pages[-1].strip() else pages
    return '\n\n'.join(f'--- page {i} ---\n\n{p.rstrip()}' for i, p in enumerate(pages, 1))


def cell_text(value, formula, fmt):
    if value is None:   # a formula without a cached value (a file written by a script): show the formula
        return formula if isinstance(formula, str) and formula.startswith('=') else ''
    if isinstance(value, datetime.datetime) and value.time() == datetime.time():
        value = value.date()
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if '%' in fmt:
            return f'{value * 100:g}%'
        if re.fullmatch(r'0+', fmt):   # zero-padded numbers such as the last four digits of a phone
            return str(int(value)).zfill(len(fmt))
        return str(int(value)) if float(value).is_integer() else repr(value)
    return str(value)


def excel_sheets(src):
    import openpyxl
    values, formulas = openpyxl.load_workbook(src, data_only=True), openpyxl.load_workbook(src)
    for ws in values.worksheets:
        fws = formulas[ws.title]
        rows = [[cell_text(c.value, fws[c.coordinate].value, c.number_format or '') for c in row] for row in ws.iter_rows()]
        rows = [r for r in rows if any(r)]
        width = max((max(i + 1 for i, v in enumerate(r) if v) for r in rows), default=0)
        yield ws.title, [r[:width] for r in rows]


def markdown_table(rows):
    if not rows:
        return '(empty sheet)'
    esc = lambda v: v.replace('|', '\\|').replace('\n', '<br>')
    lines = ['| ' + ' | '.join(esc(v) for v in r) + ' |' for r in rows]
    return '\n'.join(lines[:1] + ['|' + ' --- |' * len(rows[0])] + lines[1:])


def slug(name):
    return re.sub(r'[^\w.-]+', '_', name).strip('_') or 'sheet'


def source_root(src, work):
    src = Path(src).resolve()
    if src.is_dir():
        return src
    with zipfile.ZipFile(src) as z:   # extractall drops absolute paths and '..' parts
        z.extractall(work)
    tops = [p for p in work.iterdir() if not hidden(p.name)]
    return tops[0] if len(tops) == 1 and tops[0].is_dir() else work


def hidden(name):
    return name.startswith('.') or name == '__MACOSX'


def convert(src, out):
    """Write the bundle for case folder or zip `src` into `out` (emptied first). Returns the manifest list."""
    out = Path(out).resolve()
    shutil.rmtree(out, ignore_errors=True)
    (out / 'files').mkdir(parents=True)
    with tempfile.TemporaryDirectory() as tmp:
        root = source_root(src, Path(tmp) / 'zip')
        sources = sorted(p for p in root.rglob('*') if p.is_file() and not any(hidden(x) for x in p.relative_to(root).parts))
        entries = [convert_one(p, p.relative_to(root), out, Path(tmp)) for p in sources]
    order(entries, out)
    (out / 'manifest.json').write_text(json.dumps(entries, indent=1, ensure_ascii=False))
    write_index(entries, src, out)
    return entries


def convert_one(path, rel, out, tmp):
    entry = {'source': rel.as_posix(), 'kind': 'other', 'text': [], 'images': [], 'files': ['files/' + rel.as_posix()], 'error': None}
    original = out / 'files' / rel
    original.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, original)
    text, ext = out / 'text' / rel, path.suffix.lower()
    text.parent.mkdir(parents=True, exist_ok=True)
    try:
        if ext in LEGACY:   # LibreOffice turns .doc/.xls/.ppt into the modern format, then it converts as usual
            run('soffice', '--headless', '--convert-to', LEGACY[ext], '--outdir', str(tmp), str(path))
            path, ext = tmp / f'{path.stem}.{LEGACY[ext]}', '.' + LEGACY[ext]
        if ext in IMAGE:
            entry['kind'], entry['images'] = 'image', ['files/' + rel.as_posix()]
        elif ext in WORD:
            entry['kind'] = 'word'
            media = text.parent / (text.name + '.media')
            write(text.with_name(text.name + '.md'), word_text(path, media), entry, out)
            entry['images'] = sorted(p.relative_to(out).as_posix() for p in media.rglob('*') if p.is_file())
        elif ext == '.pdf':
            entry['kind'], body = 'pdf', pdf_text(path)
            write(text.with_name(text.name + '.md'), body, entry, out)
            if len(re.sub(r'--- page \d+ ---|\s', '', body)) < 20:
                entry['error'] = 'almost no text (a scan or a drawing): open the original with Read'
        elif ext in EXCEL:
            entry['kind'], parts = 'excel', []
            for sheet, rows in excel_sheets(path):
                parts.append(f'## Sheet: {sheet}\n\n{markdown_table(rows)}')
                buf = io.StringIO()
                csv.writer(buf).writerows(rows)
                write(text.with_name(f'{text.name}.{slug(sheet)}.csv'), buf.getvalue(), entry, out)
            write(text.with_name(text.name + '.md'), f'# {rel.name}\n\n' + '\n\n'.join(parts), entry, out, first=True)
        else:   # Markdown, CSV, JSON, YAML and any other file that decodes as text are read as they are
            entry['kind'] = 'text' if ext in TEXT else 'other text'
            write(text, decode(path.read_bytes(), ext in TEXT), entry, out)
    except UnicodeDecodeError:
        entry['kind'] = 'binary'
    except Exception as e:
        entry['error'] = f'{type(e).__name__}: {e}'[:500]
    return entry


def write(path, body, entry, out, first=False):
    path.write_text(body.rstrip() + '\n', encoding='utf-8')
    rel = path.relative_to(out).as_posix()
    entry['text'].insert(0, rel) if first else entry['text'].append(rel)


def order(entries, out):
    """BRIEF.md first, README.md second, then in the order README.md mentions them; each gets README's line."""
    readme = next((e for e in sorted(entries, key=depth) if Path(e['source']).name.lower() == 'readme.md'), None)
    lines = (out / readme['text'][0]).read_text().splitlines() if readme and readme['text'] else []
    for e in entries:
        name, stem = Path(e['source']).name.lower(), Path(e['source']).stem.lower()
        hits = [i for i, l in enumerate(lines) if name in l.lower()] or [i for i, l in enumerate(lines) if len(stem) >= 4 and stem in l.lower()]
        e['readme'] = ' '.join(lines[hits[0]].strip(' \t-*#|').split())[:300] if hits and e is not readme else ''
        e['_pos'] = hits[0] if hits else len(lines)
    brief = next((e for e in sorted(entries, key=depth) if Path(e['source']).name.lower() == 'brief.md'), None)
    entries.sort(key=lambda e: (e is not brief, e is not readme, e['_pos'], e['source']))
    for e in entries:
        del e['_pos']


def depth(e):
    return (e['source'].count('/'), e['source'])


def write_index(entries, src, out):
    cell = lambda v: v.replace('|', '\\|')
    rows = []
    for i, e in enumerate(entries, 1):
        read = ', '.join(e['text']) or ('(image: open it with Read)' if e['kind'] == 'image' else '(no text: see the original)')
        rows.append(f"| {i} | {cell(e['source'])} | {cell(read)} | {e['files'][0]} | {cell(e['readme'])} |")
    images = [f"- {img}" + (f" (inside {e['source']})" if e['kind'] != 'image' else '') for e in entries for img in e['images']]
    problems = [f"- {e['source']}: {e['error']}. The original is {e['files'][0]}." for e in entries if e['error']]
    problems += [f"- {e['source']}: binary file with no converter; the original is {e['files'][0]}." for e in entries if e['kind'] == 'binary']
    (out / 'INDEX.md').write_text('\n'.join([
        '# Case bundle', '',
        f'Converted from `{Path(src).name}` by tools/case_folder.py. Paths are relative to this folder. Read every document in this order:',
        'BRIEF.md is the starting point and README.md explains each document. `text/` holds the readable versions;',
        '`files/` holds the originals. Open every image, and any PDF whose drawings matter, with the Read tool.', '',
        '| # | Document | Read | Original | What README.md says about it |', '|---|---|---|---|---|', *rows, '',
        '## Images', '', *(images or ['(none)']), '',
        '## Problems', '', *(problems or ['(none)']), '']), encoding='utf-8')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    result = convert(sys.argv[1], sys.argv[2])
    print(f"{len(result)} documents -> {Path(sys.argv[2]) / 'INDEX.md'}; problems: {sum(bool(e['error']) or e['kind'] == 'binary' for e in result)}")
